
"""Extract text from supported local document files."""

from pathlib import Path

from docx import Document
from pypdf import PdfReader

MAX_DOCUMENT_SIZE = 20 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_document_text(file_path: str | Path) -> str:
    """Extract text from a PDF or DOCX file.

    Raises:
        FileNotFoundError: If the source file does not exist.
        ValueError: If the file is unsupported, too large, unreadable,
            encrypted, or contains no extractable text.
    """
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"Document not found: {path}")

    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError("Only PDF and DOCX files are supported.")

    if path.stat().st_size > MAX_DOCUMENT_SIZE:
        raise ValueError("Document exceeds the 20 MiB size limit.")

    try:
        if path.suffix.lower() == ".pdf":
            text = _extract_pdf_text(path)
        else:
            text = _extract_docx_text(path)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Unable to read document contents.") from exc

    text = text.strip()
    if not text:
        raise ValueError("Document contains no extractable text.")

    return text


def _extract_pdf_text(path: Path) -> str:
    """Extract text from all pages of a PDF."""
    reader = PdfReader(str(path))

    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                raise ValueError("Encrypted PDF is not supported.")
        except Exception as exc:
            raise ValueError("Encrypted PDF is not supported.") from exc

    page_texts = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(page_texts)


def _extract_docx_text(path: Path) -> str:
    """Extract paragraphs and table cell text from a DOCX file."""
    document = Document(str(path))
    parts = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)