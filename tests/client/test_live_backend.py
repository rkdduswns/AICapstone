"""Qt Client를 가상 HTTP 응답 대신 실제 FastAPI·추적기·SQLite에 연결한다."""

import socket
import threading
import time
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import httpx
import pytest
import uvicorn
from docx import Document
from PySide6.QtWidgets import QApplication

from backend.activity_store import ActivityStore
from backend.activity_tracker import ActivityTracker
from backend.app import create_app
from client.window import MainWindow


@pytest.fixture(scope="module")
def app():
    """Qt 이벤트 루프를 모듈 동안 유지한다."""
    instance = QApplication.instance() or QApplication([])
    instance.setQuitOnLastWindowClosed(False)
    return instance


def wait_for(app, condition, timeout=8):
    """실제 HTTP 콜백이 진행되도록 이벤트를 처리하며 조건을 기다린다."""
    deadline = time.monotonic() + timeout
    while not condition():
        app.processEvents()
        if time.monotonic() >= deadline:
            pytest.fail("실제 Backend 응답을 기다리는 중 시간 초과")
        time.sleep(0.005)
    app.processEvents()


@pytest.fixture
def live_backend(tmp_path):
    """운영체제 감지만 가상화하고 실제 Uvicorn 서버를 예약한 loopback 소켓에서 연다."""
    state = SimpleNamespace(process="chrome.exe", title="가상 조사 자료", tick=0.0)
    detector = SimpleNamespace(
        get_active_window_handle=lambda: 1, get_process_id=lambda handle: 42,
        get_process_name=lambda pid: state.process, normalize_app_name=lambda name: name,
        get_window_title=lambda handle: state.title,
    )
    store = ActivityStore(tmp_path / "activity.db")
    tracker = ActivityTracker(
        detector=detector, store=store, clock=lambda: state.tick,
        now=lambda: datetime(2026, 10, 21, tzinfo=timezone.utc) + timedelta(seconds=state.tick),
    )
    backend = create_app(activity_tracker=tracker, activity_store=store, collection_interval=None,
                         clipboard_reader=lambda: "가상 복사 원문", browser_token="live-test-token")
    server = uvicorn.Server(uvicorn.Config(backend, log_level="warning", access_log=False))
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
        thread = threading.Thread(target=server.run, kwargs={"sockets": [sock]}, daemon=True)
        thread.start()
        deadline = time.monotonic() + 5
        while not server.started:
            assert thread.is_alive(), "Backend가 시작 전에 종료됐습니다."
            assert time.monotonic() < deadline
            time.sleep(0.005)
        try:
            yield port, state, store
        finally:
            server.should_exit = True
            thread.join(timeout=5)
            assert not thread.is_alive()


@pytest.fixture
def window(app, live_backend):
    """서버보다 먼저 창과 모든 요청을 종료해 늦은 응답이 다음 테스트에 섞이지 않게 한다."""
    widget = MainWindow(live_backend[0])
    widget.activity.poll_timer.setInterval(20)
    widget.privacy.client.poll_timer.setInterval(20)
    widget.show()
    try:
        widget.check()
        wait_for(app, lambda: widget.privacy.state is not None and widget.application.text() != "—")
        yield widget
    finally:
        widget.close()
        app.processEvents()


def test_real_activity_pause_resume_and_program_block(app, window):
    """UI의 버튼과 설정 입력이 실제 보호 판정에 반영되고 새 상태를 Client가 수신한다."""
    assert window.application.text() == "chrome.exe"
    window.privacy.toggle.click()
    wait_for(app, lambda: "일시정지" in window.collection_status.text())
    assert window.application.text() == "—"
    wait_for(app, lambda: window.privacy.toggle.isEnabled())
    window.privacy.toggle.click()
    wait_for(app, lambda: window.application.text() == "chrome.exe")
    window.privacy.editor.setPlainText("chrome.exe")
    wait_for(app, lambda: window.privacy.apply_button.isEnabled())
    window.privacy.apply_button.click()
    wait_for(app, lambda: "수집 차단" in window.collection_status.text())
    assert window.application.text() == "—"


def test_recent_records_and_clipboard_read_real_storage(app, window, live_backend):
    """실제 수집→창 전환→저장→Qt 조회까지 원문과 작업 ID가 보존되는지 확인한다."""
    _, state, store = live_backend
    window.privacy.client.set_clipboard_enabled(True)
    wait_for(app, lambda: window.privacy.state is not None
             and window.privacy.state.settings.clipboard_enabled is True)
    wait_for(app, lambda: len(store.get_recent_clipboard_records()) == 1)
    state.process, state.tick = "word.exe", 2.0
    wait_for(app, lambda: window.application.text() == "word.exe")
    received = []
    window.recent.client.records_received.connect(received.append)
    window.tabs.setCurrentWidget(window.recent)
    wait_for(app, lambda: bool(received))
    record = next(row for row in received[0] if row.process_name == "chrome.exe")
    assert record.active_duration_ms == 2000
    # 실제 저장 ID로 별도의 Qt 클라이언트에서도 같은 복사 내용을 조회한다.
    from client.clipboard_client import ClipboardClient

    client = ClipboardClient(live_backend[0], window)
    entries = []
    client.entries_received.connect(entries.append)
    client.fetch(record.id)
    wait_for(app, lambda: bool(entries))
    assert entries[0].entries[0].text == "가상 복사 원문"
    client.close()


def test_upload_docx_and_reload_real_document_list(app, window, tmp_path):
    """실제 Qt multipart→Backend 추출→DB 저장→목록/본문 표시를 확인한다."""
    path = tmp_path / "가상 문서.docx"
    document = Document()
    document.add_paragraph("가상 통합 검증 <b>일반 텍스트</b>")
    document.save(path)
    window.tabs.setCurrentWidget(window.documents)
    wait_for(app, lambda: window.documents.loaded and not window.documents.busy)
    window.documents.upload(str(path))
    wait_for(app, lambda: len(window.documents.rows) == 1)
    row = window.documents.rows[0]
    assert row.file_name == path.name
    assert row.file_type == "docx"
    assert "<b>일반 텍스트</b>" in window.documents.content.toPlainText()
    window.documents.refresh()
    wait_for(app, lambda: not window.documents.busy)
    assert len(window.documents.rows) == 1


def test_document_import_denied_while_paused(app, window, tmp_path):
    """전체 일시정지 정책이 문서 가져오기 경로에도 적용되는지 실제 HTTP로 확인한다."""
    path = tmp_path / "가상.docx"
    document = Document()
    document.add_paragraph("허용 전 저장하지 않는 가상 자료")
    document.save(path)
    window.privacy.client.set_paused(True)
    wait_for(app, lambda: window.privacy.state is not None and window.privacy.state.settings.paused)
    window.tabs.setCurrentWidget(window.documents)
    wait_for(app, lambda: window.documents.loaded and not window.documents.busy)
    window.documents.upload(str(path))
    wait_for(app, lambda: not window.documents.busy)
    assert "허용되지" in window.documents.notice.text()
    assert window.documents.rows == []


def test_manual_document_link_reaches_recent_detail(app, window, tmp_path):
    """작업 선택→문서 연결→실제 업로드→최근 기록의 저장 본문 조회를 끝까지 실행한다."""
    window.tabs.setCurrentWidget(window.recent)
    wait_for(app, lambda: bool(window.recent.records))
    record_id = window.recent.records[0].id
    window.recent.table.selectRow(0)
    window.recent.attach_button.click()
    assert window.tabs.currentWidget() is window.documents
    assert window.documents.target_record_id == record_id
    wait_for(app, lambda: window.documents.loaded and not window.documents.busy)
    path = tmp_path / "연결 검증.docx"
    document = Document()
    document.add_paragraph("작업에 연결한 가상 <b>원문</b>")
    document.save(path)
    # Linux에서는 파일 바이트와 Windows 선택 경로만 분리 주입하고 HTTP/추출/저장은 실제 사용한다.
    source_path = r"C:\가상 자료\연결 검증.docx"
    window.documents.upload(str(path), source_path=source_path)
    wait_for(app, lambda: bool(window.documents.rows) and not window.documents.busy)
    assert window.documents.rows[0].record_id == record_id
    assert window.documents.rows[0].file_path == source_path
    wait_for(app, lambda: any(row.id == record_id and row.document is not None
                             for row in window.recent.records))
    window.tabs.setCurrentWidget(window.recent)
    index = next(i for i, row in enumerate(window.recent.records) if row.id == record_id)
    window.recent.table.selectRow(index)
    window.recent.detail_tabs.setCurrentWidget(window.recent.document_content)
    wait_for(app, lambda: window.recent.document_content.text.toPlainText()
             == "작업에 연결한 가상 <b>원문</b>")
    assert source_path in window.recent.document_content.source.toPlainText()


def test_browser_capture_reaches_recent_web_detail(app, window, live_backend):
    """실제 브라우저 수신 HTTP에 보낸 출처와 본문을 Qt 상세 화면까지 같은 ID로 조회한다."""
    with httpx.Client(base_url=f"http://127.0.0.1:{live_backend[0]}", trust_env=False) as client:
        rid = client.get("/activity/current").json()["data"]["record_id"]
        headers = {"Authorization": "Bearer live-test-token"}
        response = client.post("/browser/source", headers=headers,
            json={"browser": "chrome", "record_id": rid, "page_title": "가상 조사 자료",
                  "url": "https://example.test/research?private=query"})
        assert response.status_code == 200
        sid = response.json()["data"]["id"]
        assert client.post("/browser/content", headers=headers,
            json={"browser_source_id": sid, "content_text": "저장한 가상 웹 <b>본문</b>"}
        ).status_code == 200
    window.tabs.setCurrentWidget(window.recent)
    wait_for(app, lambda: any(row.id == rid and row.browser is not None
                             for row in window.recent.records))
    index = next(i for i, row in enumerate(window.recent.records) if row.id == rid)
    window.recent.table.selectRow(index)
    window.recent.detail_tabs.setCurrentWidget(window.recent.web_content)
    wait_for(app, lambda: window.recent.web_content.text.toPlainText()
             == "저장한 가상 웹 <b>본문</b>")
    assert "https://example.test/research" in window.recent.web_content.source.toPlainText()
    assert "private=query" not in window.recent.web_content.source.toPlainText()
